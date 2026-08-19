"""把会议纪要导出成 md、docx、pdf。

文档分两部分：前面是模型整理的纪要正文，后面附上逐句原始记录。逐句部分由程序直接从数据库
生成，不经模型，保证与左栏笔录一字不差，也杜绝改写。

docx 按用户的成稿格式：全篇仿宋，正文小四（12pt），行距 26 磅固定，两端对齐，标题分级加粗，
页脚居中页码。pdf 由 docx 转，两种格式看起来一致。
"""
from __future__ import annotations

import logging
import re
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from . import config, minutes, store

log = logging.getLogger(__name__)

BODY_FONT = "仿宋"
BODY_SIZE = Pt(12)      # 小四
LINE_SPACING = Pt(26)
H1_SIZE = Pt(15)        # 小三
H2_SIZE = Pt(14)
H3_SIZE = Pt(12)

VERBATIM_HEADING = "附：会议逐句记录"


def verbatim_markdown(ws: config.Workspace, meeting_id: int) -> str:
    """逐句原始记录。程序生成，与左栏笔录一致。"""
    turns = store.get_turns(ws, meeting_id)
    if not turns:
        return ""
    lines = [f"## {VERBATIM_HEADING}", "",
             "本节为实时转录原文，未经改写。译文由机器生成，仅供参考。", ""]
    for t in turns:
        stamp = datetime.fromtimestamp(t["ts"]).strftime("%H:%M:%S")
        src = (t.get("src") or "").strip()
        dst = (t.get("dst") or "").strip()
        if not src:
            continue
        lines.append(f"**[{stamp}]** {src}")
        if dst and dst != src:
            lines.append(f"　　{dst}")
        lines.append("")
    return "\n".join(lines)


def full_markdown(ws: config.Workspace, meeting_id: int, minutes_md: str = "") -> str:
    body = minutes_md or minutes.load(ws, meeting_id)
    verbatim = verbatim_markdown(ws, meeting_id)
    if not body:
        body = f"# 会议逐句记录\n\n**会议时间**：{minutes.meeting_window(ws, meeting_id)}\n"
    return f"{body.rstrip()}\n\n{verbatim}" if verbatim else body


# ── docx ────────────────────────────────────────────────────────────

def _style_run(run, size: Pt, bold: bool = False) -> None:
    run.font.name = BODY_FONT
    run.font.size = size
    run.bold = bold
    # 中文字体要单独指定 eastAsia，否则 Word 里中文会回退成默认字体
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)


def _add_paragraph(document, text: str, size: Pt, bold: bool = False,
                   align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent: bool = False):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing = LINE_SPACING
    fmt.space_after = Pt(0)
    fmt.space_before = Pt(6) if bold else Pt(0)
    if indent:
        fmt.first_line_indent = Pt(24)
    # 行内 **加粗** 拆成多个 run
    for piece in re.split(r"(\*\*[^*]+\*\*)", text):
        if not piece:
            continue
        strong = piece.startswith("**") and piece.endswith("**")
        run = paragraph.add_run(piece[2:-2] if strong else piece)
        _style_run(run, size, bold or strong)
    return paragraph


def _add_page_number(document) -> None:
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            mark = OxmlElement("w:fldChar")
            mark.set(qn("w:fldCharType"), kind)
            run._element.append(mark)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._element.append(instr)
    _style_run(run, Pt(10.5))


def to_docx(markdown: str, path: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Pt(72)
    section.left_margin = section.right_margin = Pt(72)

    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            _add_paragraph(document, line[4:], H3_SIZE, bold=True,
                           align=WD_ALIGN_PARAGRAPH.LEFT)
        elif line.startswith("## "):
            _add_paragraph(document, line[3:], H2_SIZE, bold=True,
                           align=WD_ALIGN_PARAGRAPH.LEFT)
        elif line.startswith("# "):
            _add_paragraph(document, line[2:], H1_SIZE, bold=True,
                           align=WD_ALIGN_PARAGRAPH.CENTER)
        elif line.startswith(("- ", "* ")):
            _add_paragraph(document, line[2:], BODY_SIZE)
        else:
            # 逐句记录那一段按原样排，不缩进；正文段落首行缩进
            indent = not line.startswith(("**[", "　　"))
            _add_paragraph(document, line, BODY_SIZE, indent=indent)

    _add_page_number(document)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def to_pdf(docx_path: Path, out_dir: Path) -> Path:
    """用 LibreOffice 转，保证 pdf 与 docx 的排版一致。"""
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="mi-pdf-") as tmp:
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp,
             str(docx_path)],
            capture_output=True, timeout=600)
        produced = list(Path(tmp).glob("*.pdf"))
        if not produced:
            detail = result.stderr.decode(errors="replace")[:200]
            raise RuntimeError(f"转 PDF 失败：{detail or '没有生成文件'}")
        target = out_dir / produced[0].name
        target.write_bytes(produced[0].read_bytes())
    return target
